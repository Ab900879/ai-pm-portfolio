# -*- coding: utf-8 -*-
"""
简历-岗位匹配度分析器 v0（MVP）
=================================
用途：读取 resume.docx + jd.txt，调用大模型输出结构化匹配报告。

设计原则（也是给面试讲的三条取舍）：
1. 不做成"直接给一个总分" —— 总分没有行动价值，用户要的是"改哪一句"
2. 不做成"自动改简历"     —— 覆盖度只有 60%，全自动改会编造经历，风险 > 收益
3. 输出必须结构化(JSON)   —— 好让下一步能接工作流 / 前端 / 统计
"""
import json
import os
import re
import sys
from pathlib import Path

from docx import Document
from openai import OpenAI

# ---------- 1. 输入层 ----------

def read_resume(path: Path) -> str:
    """读 docx 简历：段落 + 表格一起抓（简历常用表格排版，只读段落会漏一大半）"""
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_jd(path: Path) -> str:
    return path.read_text(encoding="utf-8").strip()


# ---------- 2. 提示层（这是产品的核心资产，不是代码） ----------

SYSTEM_PROMPT = """你是一位做过 5 年 AI 方向招聘的产品面试官。
你的任务不是夸候选人，而是找出**这个具体岗位的筛人逻辑**，再逐条对照简历。

三条纪律：
1. 只依据简历里写了的字，不要推测"他可能也会"。
2. 没写的就当没有，不要因为"应该会"就判为匹配。
3. 每条差距都必须给出**简历里可以直接照抄替换的改写句**，不许给"建议突出XX能力"这种空话。
"""

USER_TEMPLATE = """【目标岗位 JD】
{jd}

【候选人简历全文】
{resume}

请输出严格 JSON（不要 markdown 代码块，不要任何解释文字），格式：
{{
  "match_score": 0-100 的整数,
  "verdict": "一句话结论，20字内",
  "jd_requirements": [
    {{"req": "从JD原文摘出的要求，不要改写", "level": "must|plus", "status": "met|partial|missing", "evidence": "简历里的原句，没有就填空字符串", "note": "差距说明，20字内"}}
  ],
  "top3_gaps": [
    {{"gap": "最致命的短板，10字内", "why": "为什么HR会在意，30字内", "fix": "简历里可以直接写上去的一句话，40字内，不许编造经历"}}
  ],
  "rewrite_suggestions": [
    {{"original": "简历原句", "rewritten": "改写后", "reason": "改了什么逻辑，20字内"}}
  ],
  "interview_risk": ["面试官最可能追问的3个点，每点15字内"]
}}

要求：
- jd_requirements 最多 10 条，优先 must
- top3_gaps 只留 3 条，按严重程度排序
- rewrite_suggestions 最多 4 条，选改动收益最大的
"""


# ---------- 3. 模型调用层 ----------

def call_llm(system: str, user: str, model: str = "deepseek-chat") -> str:
    api_key = os.environ.get("DEEPSEEK_API_KEY")
    if not api_key:
        raise RuntimeError(
            "没有找到 DEEPSEEK_API_KEY 环境变量。\n"
            "Windows 设置方法：setx DEEPSEEK_API_KEY \"sk-xxx\"，然后重开终端。"
        )
    client = OpenAI(api_key=api_key, base_url="https://api.deepseek.com")
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        temperature=0.2,  # 低温度：要稳定可复现，不要创意
    )
    return resp.choices[0].message.content


def parse_json(text: str) -> dict:
    """模型有时会包 ```json 围栏，兜一层；再不行就抓第一个大括号块。"""
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.S)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", text, flags=re.S)
        if m:
            return json.loads(m.group(0))
        raise


# ---------- 4. 渲染层 ----------

ICON = {"met": "✅", "partial": "🟡", "missing": "❌"}
LVL = {"must": "硬性", "plus": "加分"}


def render(report: dict) -> str:
    lines = []
    lines.append("=" * 56)
    lines.append(f"匹配度：{report.get('match_score', '?')}/100   {report.get('verdict', '')}")
    lines.append("=" * 56)

    lines.append("\n【JD 要求逐条对照】")
    for r in report.get("jd_requirements", []):
        icon = ICON.get(r.get("status", ""), "·")
        lines.append(f"{icon} [{LVL.get(r.get('level',''), '')}] {r.get('req','')}")
        if r.get("note"):
            lines.append(f"      → {r['note']}")

    lines.append("\n【最致命的 3 个短板】")
    for i, g in enumerate(report.get("top3_gaps", []), 1):
        lines.append(f"{i}. {g.get('gap','')}｜{g.get('why','')}")
        lines.append(f"   改写句：{g.get('fix','')}")

    lines.append("\n【句子级改写建议】")
    for s in report.get("rewrite_suggestions", []):
        lines.append(f"- 原句：{s.get('original','')}")
        lines.append(f"  改为：{s.get('rewritten','')}")
        lines.append(f"  理由：{s.get('reason','')}")

    lines.append("\n【面试会被追问】")
    for r in report.get("interview_risk", []):
        lines.append(f"- {r}")
    lines.append("")
    return "\n".join(lines)


# ---------- 5. 入口 ----------

def analyze(resume_path: str, jd_path: str, save_json: bool = True) -> dict:
    resume = read_resume(Path(resume_path))
    jd = read_jd(Path(jd_path))
    print(f"[1/3] 简历读取完成：{len(resume)} 字")
    print(f"[2/3] JD 读取完成：{len(jd)} 字，正在调用模型…")
    raw = call_llm(SYSTEM_PROMPT, USER_TEMPLATE.format(jd=jd, resume=resume))
    report = parse_json(raw)
    print("[3/3] 解析完成\n")
    print(render(report))

    if save_json:
        out = Path(resume_path).with_name("report.json")
        out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"结构化结果已存：{out}")
    return report


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("用法：python analyzer.py <简历.docx> <jd.txt>")
        sys.exit(1)
    analyze(sys.argv[1], sys.argv[2])
